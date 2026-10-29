"""Экспорт в DOCX и PDF."""

from abc import ABC, abstractmethod
from pathlib import Path
from xml.sax.saxutils import escape

from ticketgen.config import AppConfig
from ticketgen.domain.models import TicketGenError, Variant
from ticketgen.logic.generator import ticket_title

# Блок — кортеж (вид, значение): ("title", str), ("text", str),
# ("table", список строк таблицы), ("break", None).


def build_tickets(header: str, variants: list[Variant]) -> list[tuple]:
    """Блоки файла с билетами (без ответов)."""
    blocks = []
    for variant in variants:
        if blocks:
            blocks.append(("break", None))
        blocks.append(("title", ticket_title(header, variant.variant_no)))
        for number, question in enumerate(variant.questions, 1):
            blocks.append(("text", f"{number}. {question.question_text}"))
    return blocks


def build_answers(header: str, variants: list[Variant]) -> list[tuple]:
    """Блоки файла с эталонными ответами для преподавателя."""
    blocks = [("title", f"{header}\nОтветы для преподавателя")]
    for variant in variants:
        blocks.append(("text", f"Билет № {variant.variant_no}"))
        for number, question in enumerate(variant.questions, 1):
            answer = question.answer_text or "—"
            blocks.append(("text", f"   {number}. {answer}"))
    return blocks


def build_summary(header: str, variants: list[Variant]) -> list[tuple]:
    """Блоки сводной ведомости: билет, состав, балл."""
    rows = [["Билет", "Вопросы (id)", "Темы", "Балл"]]
    for variant in variants:
        rows.append(
            [
                str(variant.variant_no),
                ", ".join(str(question.id) for question in variant.questions),
                ", ".join(question.topic_name for question in variant.questions),
                str(variant.total_score),
            ]
        )
    return [("title", f"{header}\nСводная ведомость"), ("table", rows)]


class BaseExporter(ABC):
    """Общий интерфейс экспортёров."""

    extension = ""

    @abstractmethod
    def save(self, blocks: list[tuple], path: Path):
        """Сохраняет блоки в файл.

        Args:
            blocks: содержимое документа.
            path: путь к файлу без расширения.
        """


class DocxExporter(BaseExporter):
    """Сохранение в Word (.docx) через python-docx."""

    extension = ".docx"

    def save(self, blocks, path):
        """Создаёт документ Word из блоков."""
        from docx import Document
        from docx.shared import Pt

        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(14)
        for kind, value in blocks:
            if kind == "title":
                document.add_heading(value, level=1)
            elif kind == "text":
                document.add_paragraph(value)
            elif kind == "break":
                document.add_page_break()
            elif kind == "table":
                table = document.add_table(rows=0, cols=len(value[0]))
                table.style = "Table Grid"
                for row in value:
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        cell.text = text
        document.save(path.with_suffix(self.extension))


class PdfExporter(BaseExporter):
    """Сохранение в PDF через reportlab с кириллическим шрифтом."""

    extension = ".pdf"
    font_name = "TicketGenFont"

    def save(self, blocks, path):
        """Создаёт PDF из блоков."""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
        )

        self._register_font()
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = self.font_name
        story = []
        for kind, value in blocks:
            if kind == "title":
                text = escape(value).replace("\n", "<br/>")
                story += [Paragraph(text, styles["Heading2"]), Spacer(1, 8)]
            elif kind == "text":
                story.append(Paragraph(escape(value), styles["Normal"]))
            elif kind == "break":
                story.append(PageBreak())
            elif kind == "table":
                cells = [
                    [Paragraph(escape(text), styles["Normal"]) for text in row]
                    for row in value
                ]
                table = Table(cells, colWidths=[40, 130, 220, 40])
                table.setStyle([("GRID", (0, 0), (-1, -1), 0.5, "black")])
                story.append(table)
        SimpleDocTemplate(str(path.with_suffix(self.extension)), pagesize=A4).build(
            story
        )

    def _register_font(self):
        """Подключает TTF-шрифт с русскими буквами.

        Без него reportlab выводит кириллицу квадратиками.
        """
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        if self.font_name in pdfmetrics.getRegisteredFontNames():
            return
        candidates = [
            AppConfig().base_dir / "fonts" / "DejaVuSans.ttf",
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        for font_path in candidates:
            if font_path.exists():
                pdfmetrics.registerFont(TTFont(self.font_name, str(font_path)))
                return
        raise TicketGenError("Не найден шрифт с кириллицей для PDF.")


class ExporterFactory:
    """Фабрика: возвращает экспортёр по названию формата."""

    _exporters = {"docx": DocxExporter, "pdf": PdfExporter}

    @classmethod
    def create(cls, file_format: str) -> BaseExporter:
        """Создаёт экспортёр для формата «docx» или «pdf»."""
        if file_format not in cls._exporters:
            raise TicketGenError(f"Неизвестный формат: {file_format}")
        return cls._exporters[file_format]()


def export_series(
    folder: Path,
    base_name: str,
    header: str,
    variants: list[Variant],
    formats: list[str],
    with_answers: bool = True,
    with_summary: bool = True,
    progress=None,
    with_tickets: bool = True,
) -> list[Path]:
    """Сохраняет серию: билеты в каждом формате, ответы и ведомость.

    Ответы и ведомость сохраняются в первом выбранном формате.

    Returns:
        Список созданных файлов.
    """
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    jobs = []
    if with_tickets:
        jobs = [(fmt, "билеты", build_tickets(header, variants)) for fmt in formats]
    if with_answers:
        jobs.append((formats[0], "ответы", build_answers(header, variants)))
    if with_summary:
        jobs.append((formats[0], "ведомость", build_summary(header, variants)))

    created = []
    for done, (fmt, suffix, blocks) in enumerate(jobs, 1):
        exporter = ExporterFactory.create(fmt)
        path = folder / f"{base_name}_{suffix}"
        try:
            exporter.save(blocks, path)
        except PermissionError as error:
            raise TicketGenError(
                f"Нет доступа к файлу {path.name}{exporter.extension}. "
                "Возможно, он открыт в Word — закройте его."
            ) from error
        created.append(path.with_suffix(exporter.extension))
        if progress:
            progress(int(done * 100 / len(jobs)))
    return created
